import os
import re
import zipfile
from pathlib import Path

try:
    import win32com.client
except Exception:
    win32com = None

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INPUT_MD = "BaoCao_PBL5_Final_Integrated.md"
OUTPUT_DOCX = "BaoCao_PBL5_Final_Formatted.docx"
OUTPUT_PDF = "BaoCao_PBL5_Final_Formatted.pdf"
CHECKLIST = "REPORT_FORMAT_CHECKLIST.md"

CONTENT_WIDTH_CM = 16.0
LOGO_MAX_CM = 5.6
IMAGE_MAX_CM = 13.5


def set_cell_shading(cell, fill="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, width_cm=CONTENT_WIDTH_CM):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(width_cm * 567)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_table(table, font_size=11, widths=None):
    set_table_width(table)
    for r_idx, row in enumerate(table.rows):
        row.height = None
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths and c_idx < len(widths):
                set_cell_width(cell, widths[c_idx])
            for para in cell.paragraphs:
                para.paragraph_format.first_line_indent = Cm(0)
                para.paragraph_format.left_indent = Cm(0)
                para.paragraph_format.right_indent = Cm(0)
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
                para.paragraph_format.line_spacing = 1.05
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    run.font.name = "Times New Roman"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                    run.font.size = Pt(font_size)
                    run.font.bold = r_idx == 0
            if r_idx == 0:
                set_cell_shading(cell, "EDEDED")


def set_page_layout(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)


def add_page_number(paragraph, fmt="Arabic"):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    for text in ("PAGE",):
        run = paragraph.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        run._r.append(fld_begin)

        run = paragraph.add_run()
        instr = OxmlElement("w:instrText")
        instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        instr.text = f" {text} \\* {fmt} "
        run._r.append(instr)

        run = paragraph.add_run()
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_end)


def add_toc_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    r = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    r._r.append(begin)

    r = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    r._r.append(instr)

    r = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r._r.append(sep)

    r = paragraph.add_run("Mục lục sẽ được cập nhật tự động trong Microsoft Word bằng thao tác Ctrl+A, F9.")
    r.italic = True

    r = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.append(end)


def setup_styles(doc):
    styles = doc.styles

    style = styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(13)
    style.font.color.rgb = RGBColor(0, 0, 0)
    pf = style.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.3
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.first_line_indent = Cm(0.75)

    h1 = styles["Heading 1"]
    h1.font.name = "Times New Roman"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(12)
    h1.paragraph_format.first_line_indent = Cm(0)

    h2 = styles["Heading 2"]
    h2.font.name = "Times New Roman"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.first_line_indent = Cm(0)

    h3 = styles["Heading 3"]
    h3.font.name = "Times New Roman"
    h3._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    h3.font.size = Pt(13)
    h3.font.bold = True
    h3.font.italic = True
    h3.font.color.rgb = RGBColor(0, 0, 0)
    h3.paragraph_format.space_before = Pt(6)
    h3.paragraph_format.space_after = Pt(6)
    h3.paragraph_format.first_line_indent = Cm(0)

    cap = styles["Caption"]
    cap.font.name = "Times New Roman"
    cap._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    cap.font.size = Pt(12)
    cap.font.italic = True
    cap.font.color.rgb = RGBColor(0, 0, 0)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(9)
    cap.paragraph_format.first_line_indent = Cm(0)

    if "Dash List" in styles:
        dash = styles["Dash List"]
    else:
        dash = styles.add_style("Dash List", WD_STYLE_TYPE.PARAGRAPH)
    dash.base_style = styles["Normal"]
    dash.font.name = "Times New Roman"
    dash._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    dash.font.size = Pt(13)
    dash.paragraph_format.left_indent = Cm(0.65)
    dash.paragraph_format.first_line_indent = Cm(0)
    dash.paragraph_format.space_after = Pt(3)
    dash.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    dash.paragraph_format.line_spacing = 1.3


def clean_inline(text):
    text = re.sub(r"<br\s*/?>", " ", text, flags=re.I)
    text = re.sub(r"</?div[^>]*>", "", text, flags=re.I)
    text = re.sub(r"\[([^\]]+)\]\(#[^)]+\)", r"\1", text)
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    return text.strip()


def add_runs_from_markdown(paragraph, text, default_bold=False):
    text = clean_inline(text)
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
            run.bold = default_bold
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def preprocess_markdown(raw):
    raw = re.sub(r"<!--.*?-->", "", raw, flags=re.DOTALL)
    raw = re.sub(r"</?div[^>]*>", "", raw, flags=re.I)

    # Remove manual Markdown TOC completely.
    raw = re.sub(r"# MỤC LỤC\s+.*?(?=\n# DANH SÁCH HÌNH VẼ)", "# MỤC LỤC\n\n", raw, flags=re.DOTALL)

    # Remove residual internal Markdown links if any.
    raw = re.sub(r"\[([^\]]+)\]\(#[^)]+\)", r"\1", raw)
    return raw


def extract_cover_and_body(content):
    m = re.search(r"(?m)^# TÓM TẮT ĐỒ ÁN\s*$", content)
    if not m:
        return content, ""
    return content[:m.start()], content[m.start():]


def parse_table(lines, i):
    data = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        row = [clean_inline(c.strip()).replace("**", "").replace("*", "") for c in lines[i].strip().strip("|").split("|")]
        if not all(re.match(r"^:?-{3,}:?$", c) for c in row):
            data.append(row)
        i += 1
    max_cols = max((len(r) for r in data), default=0)
    for row in data:
        row.extend([""] * (max_cols - len(row)))
    return data, i


def add_md_table(doc, table_data, caption=None, list_table=False):
    if caption:
        p = doc.add_paragraph(caption, style="Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Cm(0)

    rows, cols = len(table_data), len(table_data[0]) if table_data else 0
    if rows == 0 or cols == 0:
        return

    table = doc.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"

    for r_idx, row in enumerate(table_data):
        for c_idx, val in enumerate(row):
            table.cell(r_idx, c_idx).text = val

    if caption and "Bảng 3.1" in caption:
        widths = [2.6, 3.0, 3.4, 3.4, 3.6]
        font_size = 9.5
    elif cols <= 3:
        widths = [2.2, 11.5, 2.3] if list_table and cols == 3 else None
        font_size = 11
    elif cols == 4:
        widths = [1.3, 4.0, 8.2, 2.5]
        font_size = 10.5
    elif cols == 5:
        widths = [1.2, 3.0, 3.8, 4.5, 3.5]
        font_size = 10
    else:
        widths = [CONTENT_WIDTH_CM / cols] * cols
        font_size = 9.5

    format_table(table, font_size=font_size, widths=widths)


def add_cover_page(doc, cover_md):
    section = doc.sections[0]
    set_page_layout(section)

    def centered(text="", size=13, bold=False, space_after=3):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = 1.05
        if text:
            run = p.add_run(text)
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(size)
            run.bold = bold
        return p

    centered("ĐẠI HỌC ĐÀ NẴNG", 13, True, 2)
    centered("TRƯỜNG ĐẠI HỌC BÁCH KHOA", 13, True, 2)
    centered("KHOA CÔNG NGHỆ THÔNG TIN", 13, True, 6)

    logo_match = re.search(r"!\[[^\]]*\]\(([^)]+)\)", cover_md)
    if logo_match and os.path.exists(logo_match.group(1)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run()
        try:
            run.add_picture(logo_match.group(1), width=Cm(LOGO_MAX_CM))
        except Exception:
            pass

    centered("PBL5: DỰ ÁN KỸ THUẬT MÁY TÍNH", 16, True, 7)
    centered("Đề tài: Xây dựng hệ thống Thùng rác thông minh", 14, True, 9)
    centered("GIẢNG VIÊN HƯỚNG DẪN: TS. Huỳnh Hữu Hưng", 13, True, 5)
    centered("SINH VIÊN THỰC HIỆN", 13, True, 3)

    students = [
        ["Họ và Tên", "Lớp", "MSSV"],
        ["Nguyễn Quốc Nguyên", "23T_DT4", "102230362"],
        ["Đặng Đăng Khoa", "23T_DT4", "102230352"],
        ["Lê Bá Nguyên Long", "23T_DT4", "102230357"],
        ["Phan Thị Nhân Vỹ", "23T_DT1", "102230225"],
    ]
    table = doc.add_table(rows=len(students), cols=3)
    table.style = "Table Grid"
    for r, row in enumerate(students):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    format_table(table, font_size=11.5, widths=[7.8, 3.8, 4.4])

    for _ in range(2):
        centered("", 8, False, 0)
    p = centered("Đà Nẵng, 06/2026", 13, True, 0)
    p.paragraph_format.space_before = Pt(12)


def add_section(doc, start=WD_SECTION_START.NEW_PAGE, roman=False, start_num=None, page_numbers=True):
    section = doc.add_section(start)
    set_page_layout(section)
    section.footer.is_linked_to_previous = False
    section.header.is_linked_to_previous = False
    footer = section.footer
    for p in footer.paragraphs:
        p.clear()
    if page_numbers:
        add_page_number(footer.paragraphs[0], "roman" if roman else "Arabic")
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:fmt"), "lowerRoman" if roman else "decimal")
    if start_num is not None:
        pg_num.set(qn("w:start"), str(start_num))
    return section


def build_docx(md_path, out_path):
    warnings = []
    raw = Path(md_path).read_text(encoding="utf-8")
    content = preprocess_markdown(raw)
    cover_md, body_md = extract_cover_and_body(content)

    code_blocks = []

    def code_replacer(match):
        code_blocks.append(match.group(0))
        return f"\n[CODE_BLOCK_{len(code_blocks) - 1}]\n"

    body_md = re.sub(r"```.*?```", code_replacer, body_md, flags=re.DOTALL)

    doc = Document()
    setup_styles(doc)
    add_cover_page(doc, cover_md)

    # Front matter section starts after cover.
    add_section(doc, WD_SECTION_START.NEW_PAGE, roman=True, start_num=1, page_numbers=True)

    lines = body_md.splitlines()
    pending_table_caption = None
    pending_image_caption = None
    i = 0
    while i < len(lines):
        raw_line = lines[i]
        line = raw_line.strip()

        if not line:
            i += 1
            continue
        if line == "---":
            i += 1
            continue

        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = clean_inline(line.lstrip("#").strip())
            if text == "CHƯƠNG 1: GIỚI THIỆU":
                add_section(doc, WD_SECTION_START.NEW_PAGE, roman=False, start_num=1, page_numbers=True)
            elif text in ("TÀI LIỆU THAM KHẢO", "PHỤ LỤC: MÃ NGUỒN"):
                section = add_section(doc, WD_SECTION_START.NEW_PAGE, roman=False, start_num=None, page_numbers=True)
                pg_num = section._sectPr.find(qn("w:pgNumType"))
                if pg_num is not None:
                    section._sectPr.remove(pg_num)

            if text == "MỤC LỤC":
                p = doc.add_paragraph(text, style="Heading 1")
                p = doc.add_paragraph()
                add_toc_field(p)
                i += 1
                continue

            style = "Heading 1" if level == 1 else "Heading 2" if level == 2 else "Heading 3" if level == 3 else "Normal"
            doc.add_paragraph(text.upper() if level == 1 else text, style=style)
            i += 1
            continue

        if line.startswith("**Bảng") and line.endswith("**"):
            pending_table_caption = line.replace("**", "")
            i += 1
            continue

        if line.startswith("**Hình") and line.endswith("**"):
            pending_image_caption = line.replace("**", "")
            i += 1
            continue

        img = re.match(r"!\[(.*?)\]\((.*?)\)", line)
        if img:
            alt_text, img_path = img.group(1), img.group(2)
            if os.path.exists(img_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
                try:
                    p.add_run().add_picture(img_path, width=Cm(IMAGE_MAX_CM))
                except Exception as e:
                    warnings.append(f"Không chèn được hình {img_path}: {e}")
            else:
                warnings.append(f"Thiếu hình: {img_path}")
                p = doc.add_paragraph(f"[Thiếu hình: {alt_text} - {img_path}]", style="Caption")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Caption must be below image. Prefer next line if it is caption.
            caption = pending_image_caption
            if i + 1 < len(lines) and lines[i + 1].strip().startswith("**Hình"):
                caption = lines[i + 1].strip().replace("**", "")
                i += 1
            if caption:
                cap = doc.add_paragraph(caption, style="Caption")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pending_image_caption = None
            i += 1
            continue

        if line.startswith("|"):
            table_data, i = parse_table(lines, i)
            prev_heading = ""
            for p in reversed(doc.paragraphs[-4:]):
                if p.style.name.startswith("Heading"):
                    prev_heading = p.text
                    break
            list_table = prev_heading in ("DANH SÁCH HÌNH VẼ", "DANH SÁCH BẢNG")
            add_md_table(doc, table_data, pending_table_caption, list_table=list_table)
            pending_table_caption = None
            continue

        if line.startswith("- ") or line.startswith("* "):
            items = []
            while i < len(lines) and lines[i].strip().startswith(("- ", "* ")):
                item = clean_inline(lines[i].strip()[2:].strip())
                item = item.rstrip(";.")
                item = re.sub(r":;$", ";", item)
                item = re.sub(r":\.$", ".", item)
                item = re.sub(r":$", "", item)
                items.append(item)
                i += 1
            for idx, item in enumerate(items):
                suffix = "." if idx == len(items) - 1 else ";"
                p = doc.add_paragraph(style="Dash List")
                p.paragraph_format.first_line_indent = Cm(0)
                add_runs_from_markdown(p, f"- {item}{suffix}")
            continue

        code_match = re.match(r"\[CODE_BLOCK_(\d+)\]", line)
        if code_match:
            p = doc.add_paragraph("Chi tiết mã nguồn được trình bày tại Phụ lục.", style="Normal")
            p.runs[0].italic = True
            i += 1
            continue

        p = doc.add_paragraph(style="Normal")
        add_runs_from_markdown(p, line)
        i += 1

    if code_blocks:
        add_section(doc, WD_SECTION_START.NEW_PAGE, roman=False, start_num=None, page_numbers=True)
        doc.add_paragraph("PHỤ LỤC: MÃ NGUỒN", style="Heading 1")
        for idx, code in enumerate(code_blocks, start=1):
            code_text = re.sub(r"^```[^\n]*\n?", "", code.strip())
            code_text = re.sub(r"\n?```$", "", code_text)
            doc.add_paragraph(f"Phụ lục A.{idx}", style="Heading 2")
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(code_text)
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(9)

    doc.save(out_path)
    return warnings


def try_word_com_update(docx_path, pdf_path):
    return False, "Bỏ qua update bằng Word COM để tránh treo tiến trình. Người dùng tự update TOC bằng Ctrl+A -> F9."


def inspect_docx(docx_path):
    result = {
        "exists": os.path.exists(docx_path),
        "size": os.path.getsize(docx_path) if os.path.exists(docx_path) else 0,
        "has_div": False,
        "has_md_anchor_link": False,
        "has_round_bullet": False,
        "section_count": 0,
        "paragraph_count": 0,
        "table_count": 0,
        "image_count": 0,
        "toc_field": False,
        "page_estimate": None,
        "errors": [],
    }
    if not result["exists"]:
        result["errors"].append("File DOCX không tồn tại.")
        return result

    doc = Document(docx_path)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text += "\n" + cell.text

    result["has_div"] = "<div" in full_text or "</div>" in full_text
    result["has_md_anchor_link"] = "](#" in full_text
    result["has_round_bullet"] = "" in full_text or "•" in full_text
    result["section_count"] = len(doc.sections)
    result["paragraph_count"] = len(doc.paragraphs)
    result["table_count"] = len(doc.tables)

    with zipfile.ZipFile(docx_path) as zf:
        xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
        result["toc_field"] = 'TOC \\o "1-3"' in xml or " TOC " in xml or 'TOC \\\\o "1-3"' in xml
        result["image_count"] = len([n for n in zf.namelist() if n.startswith("word/media/")])
        explicit_breaks = xml.count('w:type w:val="page"') + xml.count("lastRenderedPageBreak")
        result["page_estimate"] = max(1, explicit_breaks + 1)

    if result["has_div"]:
        result["errors"].append("Còn tag <div> hoặc </div> trong nội dung DOCX.")
    if result["has_md_anchor_link"]:
        result["errors"].append("Còn Markdown link kiểu ](#anchor).")
    if result["has_round_bullet"]:
        result["errors"].append("Còn bullet tròn/ trong DOCX.")
    if result["section_count"] < 3:
        result["errors"].append("Số section thấp hơn yêu cầu tối thiểu.")
    return result


def write_checklist(path, warnings, word_ok, word_msg, inspection):
    status = "ĐẠT" if not inspection["errors"] else "CHƯA ĐẠT"
    content = f"""# REPORT_FORMAT_CHECKLIST

## 1. File đầu ra
- DOCX: `{OUTPUT_DOCX}`
- PDF: `{OUTPUT_PDF}` {"(đã xuất)" if os.path.exists(OUTPUT_PDF) else "(chưa xuất được tự động)"}
- Trạng thái kiểm tra tự động: **{status}**

## 2. Đã sửa theo phản hồi
- [x] Fix mục lục: "TÀI LIỆU THAM KHẢO" không hiển thị trang 1, tiếp tục đánh số trang theo section trước.
- [x] Fix bảng 3.1: Giảm font xuống 9.5, tăng độ rộng cột Lớp để không vỡ chữ (Hazardous, Biological...).
- [x] Sửa lỗi dấu câu API: loại bỏ các trường hợp như `:;`, `:` ở cuối dòng trong list.
- [x] Tiền xử lý Markdown: xóa HTML wrapper `<div align="center">`, `</div>` và comment blocks.
- [x] Xóa mục lục Markdown thủ công từ `# MỤC LỤC` đến trước `# DANH SÁCH HÌNH VẼ`.
- [x] Xóa Markdown anchor links kiểu `[text](#anchor)` khỏi nội dung DOCX.
- [x] Tạo trang bìa riêng bằng `python-docx`, căn giữa, logo giới hạn khoảng {LOGO_MAX_CM} cm, bảng sinh viên fit một trang.
- [x] Tạo Section Break Next Page sau trang bìa và trước Chương 1; thêm section riêng cho Tài liệu tham khảo/Phụ lục nếu có.
- [x] Thiết lập A4, lề Trên 2 cm, Dưới 2 cm, Trái 3 cm, Phải 2 cm.
- [x] Thiết lập Styles: Normal, Heading 1, Heading 2, Heading 3, Caption.
- [x] Bullet được chuyển thành đoạn văn bắt đầu bằng dấu gạch ngang `-`, không dùng bullet tròn.
- [x] Hình nội dung căn giữa, không kéo méo ảnh, giới hạn chiều rộng tối đa {IMAGE_MAX_CM} cm; caption hình đặt dưới hình.
- [x] Caption bảng đặt trên bảng khi Markdown có caption dạng `**Bảng ...**`.
- [x] Bảng nhiều cột dùng font nhỏ hơn và fixed layout để hạn chế chữ bị tách dòng xấu.
- [x] Code block dài được chuyển sang Phụ lục.

## 3. Word COM / Field động
- Word COM: {"Đã chạy thành công" if word_ok else "Không chạy được hoặc bị lỗi"}
- Ghi chú: {word_msg}
- Nếu mục lục/số trang chưa hiển thị đúng trong Word, mở file DOCX rồi nhấn **Ctrl+A → F9** và chọn cập nhật toàn bộ bảng.

## 4. Kết quả kiểm tra tự động
- File tồn tại: {inspection["exists"]}
- Dung lượng: {inspection["size"]:,} bytes
- Không còn `<div` hoặc `</div>`: {not inspection["has_div"]}
- Không còn Markdown link `](#`: {not inspection["has_md_anchor_link"]}
- Không còn bullet tròn/``: {not inspection["has_round_bullet"]}
- Số section: {inspection["section_count"]}
- Số đoạn: {inspection["paragraph_count"]}
- Số bảng: {inspection["table_count"]}
- Số ảnh nhúng: {inspection["image_count"]}
- Có TOC field/placeholder sạch: {inspection["toc_field"]}
- Ước lượng số trang thô từ XML: {inspection["page_estimate"]}

## 5. Cảnh báo
"""
    if warnings:
        for w in warnings:
            content += f"- {w}\n"
    else:
        content += "- Không phát hiện thiếu ảnh khi chèn bằng python-docx.\n"

    content += "\n## 6. Lỗi còn lại nếu có\n"
    if inspection["errors"]:
        for e in inspection["errors"]:
            content += f"- {e}\n"
    else:
        content += "- Không có lỗi trong các tiêu chí kiểm tra tự động bắt buộc.\n"

    Path(path).write_text(content, encoding="utf-8")


if __name__ == "__main__":
    print("Building DOCX v3 from cleaned Markdown...")
    warnings = build_docx(INPUT_MD, OUTPUT_DOCX)

    print("Trying Word COM field update/PDF export...")
    word_ok, word_msg = try_word_com_update(OUTPUT_DOCX, OUTPUT_PDF)
    print(word_msg)

    print("Inspecting output DOCX...")
    inspection = inspect_docx(OUTPUT_DOCX)
    for key, value in inspection.items():
        if key != "errors":
            print(f"{key}: {value}")
    if inspection["errors"]:
        print("VALIDATION FAILED:")
        for err in inspection["errors"]:
            print(f"- {err}")
    else:
        print("VALIDATION PASSED")

    write_checklist(CHECKLIST, warnings, word_ok, word_msg, inspection)
    print(f"Wrote {CHECKLIST}")

    if inspection["errors"]:
        raise SystemExit(1)
